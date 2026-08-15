"""Unit tests for the Excel generation layer.

Covers ExcelGeneratorService (generation, validation, file-name normalisation)
and GenerateExcelHandler (ActionResult adaptation).
"""

import io

import pytest
from openpyxl import load_workbook

from app.execution.python.base.exceptions import (
    EmptyPayloadError,
    InvalidConfigurationError,
)
from app.execution.python.base.generate_model import GeneratedFile
from app.execution.python.excel.generator import (
    XLSX_MIME_TYPE,
    ExcelGeneratorService,
)
from app.execution.python.excel.handler import GenerateExcelHandler
from app.workflow_execution.schemas.action_result import ActionResult


@pytest.fixture
def payload() -> dict:
    return {
        "customer_name": "Asha Rao",
        "loan_number": "LN-2024-0001",
        "branch": "Bengaluru MG Road",
    }


@pytest.fixture
def config() -> dict:
    return {"headers": ["customer_name", "loan_number", "branch"]}


@pytest.fixture
def storage(tmp_path):
    """A FileStorageService backed by an isolated temp directory."""
    from app.storage.local import LocalStorageProvider
    from app.storage.service import FileStorageService

    return FileStorageService(LocalStorageProvider(base_path=tmp_path))


def _read_rows(content: bytes) -> list[list]:
    """Load .xlsx bytes back into rows for assertions."""
    workbook = load_workbook(io.BytesIO(content))
    worksheet = workbook.active
    return [list(row) for row in worksheet.iter_rows(values_only=True)]


# ── ExcelGeneratorService ─────────────────────────────────────────────────────

def test_generate_succeeds(payload, config):
    result = ExcelGeneratorService().generate(payload, config)
    assert isinstance(result, GeneratedFile)


def test_file_name_defaults_to_document_xlsx(payload, config):
    result = ExcelGeneratorService().generate(payload, config)
    assert result.file_name == "document.xlsx"


def test_file_name_uses_configured_output_name(payload, config):
    result = ExcelGeneratorService().generate(
        payload, {**config, "output_name": "loan_export.xlsx"}
    )
    assert result.file_name == "loan_export.xlsx"


def test_xlsx_extension_is_appended_when_missing(payload, config):
    result = ExcelGeneratorService().generate(
        payload, {**config, "output_name": "loan_export"}
    )
    assert result.file_name == "loan_export.xlsx"


def test_mime_type_is_xlsx(payload, config):
    result = ExcelGeneratorService().generate(payload, config)
    assert result.mime_type == XLSX_MIME_TYPE
    assert result.extension == "xlsx"


def test_content_length_is_positive(payload, config):
    result = ExcelGeneratorService().generate(payload, config)
    assert result.size > 0
    assert result.size == len(result.content)


def test_content_starts_with_zip_magic_bytes(payload, config):
    """A real .xlsx is a ZIP container, so it begins with PK\\x03\\x04."""
    result = ExcelGeneratorService().generate(payload, config)
    assert result.content.startswith(b"PK\x03\x04")


def test_header_and_data_rows_are_written(payload, config):
    result = ExcelGeneratorService().generate(payload, config)
    rows = _read_rows(result.content)
    assert rows[0] == ["customer_name", "loan_number", "branch"]
    assert rows[1] == ["Asha Rao", "LN-2024-0001", "Bengaluru MG Road"]


def test_missing_values_are_written_as_blank_cells(config):
    result = ExcelGeneratorService().generate({"customer_name": "Asha"}, config)
    rows = _read_rows(result.content)
    # openpyxl stores an empty string as a blank cell, which reads back as None.
    assert rows[1] == ["Asha", None, None]


def test_sheet_name_is_configurable(payload, config):
    result = ExcelGeneratorService().generate(
        payload, {**config, "sheet_name": "Loans"}
    )
    workbook = load_workbook(io.BytesIO(result.content))
    assert workbook.active.title == "Loans"


def test_empty_payload_is_rejected(config):
    with pytest.raises(EmptyPayloadError):
        ExcelGeneratorService().generate({}, config)


def test_missing_headers_are_rejected(payload):
    with pytest.raises(InvalidConfigurationError):
        ExcelGeneratorService().generate(payload, {})


# ── GenerateExcelHandler ──────────────────────────────────────────────────────

def test_handler_returns_successful_action_result(payload, config, storage):
    result = GenerateExcelHandler(storage=storage)(
        payload, {**config, "output_name": "export"}
    )
    assert isinstance(result, ActionResult)
    assert result.success is True
    assert result.outputs["file_name"] == "export.xlsx"
    assert result.outputs["mime_type"] == XLSX_MIME_TYPE
    assert result.outputs["size"] > 0
    assert result.outputs["storage_provider"] == "local"
    assert result.outputs["file_id"]


def test_handler_stores_a_retrievable_xlsx(payload, config, storage):
    result = GenerateExcelHandler(storage=storage)(payload, config)
    content = storage.retrieve(result.outputs["storage_path"])
    assert content.startswith(b"PK\x03\x04")
    rows = _read_rows(content)
    assert rows[0] == ["customer_name", "loan_number", "branch"]


def test_handler_empty_payload_returns_failure(config, storage):
    result = GenerateExcelHandler(storage=storage)({}, config)
    assert result.success is False
    assert result.error is not None


def test_handler_missing_headers_returns_failure(payload, storage):
    result = GenerateExcelHandler(storage=storage)(payload, {})
    assert result.success is False
    assert result.error is not None
